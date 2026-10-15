#!/usr/bin/env python3
"""Genera i template Excel di importazione Odoo 20.0 (repo pixora/odoo, branch 20.0).

Uso:  python3 genera_template.py   (dalla cartella import_templates)
Output: file 01..06 *.xlsx nella stessa cartella.

I nomi colonna sono i nomi TECNICI dei campi Odoo: l'import li riconosce in
automatico indipendentemente dalla lingua dell'utente.
Campi Community verificati sul sorgente 20.0 del repo. Budget analitico confermato
su 20 dal modulo pixora/demo-alina; budget finanziario verificato su Odoo 19.
"""
import csv
import os

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

HERE = os.path.dirname(os.path.abspath(__file__))
L10N_IT_COA = os.path.join(HERE, '..', 'addons', 'l10n_it', 'data', 'template', 'account.account-it.csv')

RED = PatternFill('solid', fgColor='C0392B')
BLUE = PatternFill('solid', fgColor='2E5C8A')
GREY = PatternFill('solid', fgColor='7F8C8D')
EXAMPLE = PatternFill('solid', fgColor='FFF4CC')
HEAD_FONT = Font(bold=True, color='FFFFFF')
THIN = Border(*(Side(style='thin', color='BBBBBB'),) * 4)

# Spec colonna: (campo_tecnico, etichetta_IT, livello, tipo, valori/formato, note, larghezza, lista_valori)
# livello: 'O' obbligatorio, 'C' consigliato, 'F' facoltativo
ACCOUNT_TYPES = [
    ('asset_receivable', 'Crediti (clienti) - riconciliabile obbligatorio'),
    ('asset_cash', 'Banca e cassa'),
    ('asset_current', 'Attività correnti'),
    ('asset_non_current', 'Attività non correnti'),
    ('asset_prepayments', 'Risconti attivi / anticipi'),
    ('asset_fixed', 'Immobilizzazioni'),
    ('liability_payable', 'Debiti (fornitori) - riconciliabile obbligatorio'),
    ('liability_credit_card', 'Carta di credito'),
    ('liability_current', 'Passività correnti'),
    ('liability_non_current', 'Passività non correnti'),
    ('equity', 'Patrimonio netto'),
    ('equity_unaffected', 'Utile/perdita esercizio corrente (UNO solo per azienda)'),
    ('income', 'Ricavi'),
    ('income_other', 'Altri ricavi'),
    ('expense', 'Costi'),
    ('expense_other', 'Altri costi'),
    ('expense_depreciation', 'Ammortamenti'),
    ('expense_direct_cost', 'Costo del venduto'),
    ('off_balance', 'Conti d\'ordine (fuori bilancio)'),
]
IT_TAGS = [
    ('l10n_it.account_tag_A_ATT', 'SP Attivo A - Crediti verso soci'),
    ('l10n_it.account_tag_B_ATT', 'SP Attivo B - Immobilizzazioni'),
    ('l10n_it.account_tag_C_ATT', 'SP Attivo C - Attivo circolante'),
    ('l10n_it.account_tag_D_ATT', 'SP Attivo D - Ratei e risconti attivi'),
    ('l10n_it.account_tag_A_PASS', 'SP Passivo A - Patrimonio netto'),
    ('l10n_it.account_tag_B_PASS', 'SP Passivo B - Fondi rischi e oneri'),
    ('l10n_it.account_tag_C_PASS', 'SP Passivo C - TFR'),
    ('l10n_it.account_tag_D_PASS', 'SP Passivo D - Debiti'),
    ('l10n_it.account_tag_E_PASS', 'SP Passivo E - Ratei e risconti passivi'),
    ('l10n_it.account_tag_RISCHI', 'Conti d\'ordine - Rischi'),
    ('l10n_it.account_tag_IMPEGNI', 'Conti d\'ordine - Impegni'),
    ('l10n_it.account_tag_BENI', 'Conti d\'ordine - Beni di terzi'),
    ('l10n_it.account_tag_A_PL', 'CE A - Valore della produzione'),
    ('l10n_it.account_tag_B_PL', 'CE B - Costi della produzione'),
    ('l10n_it.account_tag_C_PL', 'CE C - Proventi e oneri finanziari'),
    ('l10n_it.account_tag_D_PL', 'CE D - Rettifiche di valore attività finanziarie'),
    ('l10n_it.account_tag_E_PL', 'CE E - Proventi e oneri straordinari'),
]
PAYMENT_TERMS = [
    ('account.account_payment_term_immediate', 'Pagamento immediato'),
    ('account.account_payment_term_15days', '15 giorni'),
    ('account.account_payment_term_21days', '21 giorni'),
    ('account.account_payment_term_30days', '30 giorni'),
    ('account.account_payment_term_45days', '45 giorni'),
    ('account.account_payment_term_end_following_month', 'Fine mese successivo'),
    ('account.account_payment_term_30_days_end_month_the_10', '30 gg fine mese il 10'),
    ('account.account_payment_term_90days_on_the_10th', '90 gg il 10'),
]
BOOL = ['TRUE', 'FALSE']
EMP_TYPES = [
    ('hr.contract_type_employee', 'Dipendente'),
    ('hr.contract_type_company_executive', 'Dirigente / amministratore'),
    ('hr.contract_type_interim', 'Interinale'),
    ('hr.contract_type_seasonal', 'Occasionale / stagionale'),
    ('hr.contract_type_apprenticeship', 'Apprendista'),
    ('hr.contract_type_intern', 'Stagista / tirocinante'),
    ('hr.contract_type_student', 'Studente'),
    ('hr.contract_type_thesis', 'Tesista'),
]
PROVINCE = ('Usare base.state_it_ + sigla provincia minuscola: base.state_it_mi (Milano), '
            'base.state_it_rm (Roma), base.state_it_to (Torino)...')
EXT_ID = ('Codice univoco della riga (es. FORN_0001). Odoo lo salva come __import__.FORN_0001: '
          'reimportando lo stesso file AGGIORNA invece di duplicare.')

TEMPLATES = [
    {
        'file': '01_Fornitori.xlsx',
        'model': 'res.partner',
        'title': 'Fornitori (Contatti)',
        'menu': 'Contabilità > Fornitori > Fornitori  (oppure Contatti)',
        'cols': [
            ('id', 'ID esterno', 'C', 'Testo', 'FORN_0001', EXT_ID, 14, None),
            ('is_company', 'È un\'azienda', 'C', 'Booleano', 'TRUE / FALSE', 'TRUE = società, FALSE = persona fisica/ditta individuale.', 12, BOOL),
            ('name', 'Nome', 'O', 'Testo', 'Ragione sociale', 'Unico campo realmente obbligatorio.', 32, None),
            ('ref', 'Riferimento', 'F', 'Testo', 'Codice fornitore gestionale precedente', 'Utile per riconciliare con il vecchio sistema.', 14, None),
            ('vat', 'Partita IVA', 'C', 'Testo', 'IT01234567890', 'Con prefisso paese. Odoo la valida (checksum): se errata blocca la riga.', 16, None),
            ('l10n_it_codice_fiscale', 'Codice Fiscale', 'C', 'Testo (max 16)', 'RSSMRA80A01H501U / 01234567890', 'Richiesto per FatturaPA se diverso dalla P.IVA (persone fisiche). Modulo l10n_it_edi.', 18, None),
            ('street', 'Via', 'C', 'Testo', 'Via Roma 1', 'Obbligatorio per fatturazione elettronica.', 26, None),
            ('street2', 'Via 2', 'F', 'Testo', '', '', 14, None),
            ('zip', 'CAP', 'C', 'Testo', '20100', 'Formattare la colonna come TESTO (zeri iniziali).', 8, None),
            ('city', 'Città', 'C', 'Testo', 'Milano', '', 16, None),
            ('state_id/id', 'Provincia (ID esterno)', 'C', 'ID esterno', 'base.state_it_mi', PROVINCE, 18, None),
            ('country_id/id', 'Nazione (ID esterno)', 'C', 'ID esterno', 'base.it', 'base.it, base.de, base.fr ... (base. + codice ISO minuscolo).', 12, None),
            ('phone', 'Telefono', 'F', 'Testo', '+39 02 1234567', '', 16, None),
            ('email', 'Email', 'F', 'Testo', 'amministrazione@fornitore.it', '', 26, None),
            ('website', 'Sito web', 'F', 'Testo', '', '', 18, None),
            ('l10n_it_pec_email', 'PEC', 'F', 'Testo', 'fornitore@pec.it', 'Modulo l10n_it_edi.', 22, None),
            ('l10n_it_pa_index', 'Codice destinatario SDI', 'F', 'Testo (6/7 car.)', '0000000 / ABC1234', 'Per i fornitori serve solo se emetti verso di loro (es. note di debito). 7 car. privati, 6 PA.', 14, None),
            ('supplier_rank', 'Rango fornitore', 'O', 'Intero', '1', 'METTERE 1: è ciò che fa comparire il contatto nel menu Fornitori.', 10, None),
            ('property_supplier_payment_term_id/id', 'Termini pagamento fornitore', 'F', 'ID esterno', 'account.account_payment_term_30days', 'Vedi foglio Valori. Per termini custom (es. 60 gg DF FM) usare la colonna property_supplier_payment_term_id con il NOME esatto.', 30, [p[0] for p in PAYMENT_TERMS]),
            ('property_account_payable_id', 'Conto debiti', 'F', 'Codice conto', '2501', 'Lasciare vuoto per usare il default aziendale (2501 Debiti v/fornitori).', 12, None),
            ('property_account_position_id', 'Posizione fiscale', 'F', 'Nome', 'Regime Extra comunitario', 'Nome ESATTO della posizione fiscale presente in Odoo.', 22, None),
            ('category_id', 'Etichette', 'F', 'Nomi separati da virgola', 'Materie prime,Strategico', 'Le etichette inesistenti vengono proposte per la creazione.', 22, None),
            ('lang', 'Lingua', 'F', 'Codice', 'it_IT', 'it_IT, en_US, de_DE ... (lingua deve essere attiva).', 8, ['it_IT', 'en_US', 'de_DE', 'fr_FR', 'es_ES']),
            ('comment', 'Note', 'F', 'Testo', '', '', 20, None),
        ],
        'examples': [
            ['FORN_0001', 'TRUE', 'ESEMPIO Metalli Srl', 'F001', 'IT01234567890', '01234567890', 'Via Roma 1', '', '20121', 'Milano', 'base.state_it_mi', 'base.it', '+39 02 1234567', 'amm@esempio-metalli.it', '', 'esempio-metalli@pec.it', '0000000', 1, 'account.account_payment_term_30days', '2501', '', 'Materie prime', 'it_IT', ''],
            ['FORN_0002', 'FALSE', 'ESEMPIO Mario Rossi (ditta individuale)', 'F002', 'IT09876543210', 'RSSMRA80A01H501U', 'Via Verdi 10', '', '00184', 'Roma', 'base.state_it_rm', 'base.it', '', 'mario.rossi@example.com', '', '', '', 1, 'account.account_payment_term_end_following_month', '', '', '', 'it_IT', ''],
            ['FORN_0003', 'TRUE', 'ESEMPIO Stahl GmbH', 'F003', 'DE123456789', '', 'Hauptstrasse 5', '', '80331', 'München', '', 'base.de', '', 'info@example.de', '', '', 'XXXXXXX', 1, 'account.account_payment_term_45days', '', 'Regime Intra comunitario', 'Estero', 'de_DE', ''],
        ],
    },
    {
        'file': '02_Conti_Bancari.xlsx',
        'model': 'res.partner.bank',
        'title': 'Conti bancari (IBAN + banca) di fornitori, dipendenti e azienda',
        'menu': 'Contabilità > Configurazione > Conti bancari  (Impostazioni > Tecnico se non visibile)',
        'cols': [
            ('id', 'ID esterno', 'C', 'Testo', 'IBAN_FORN_0001', EXT_ID, 16, None),
            ('partner_id/id', 'Titolare (ID esterno)', 'O', 'ID esterno', 'FORN_0001', 'ID esterno del fornitore importato nel file 01 (scrivere FORN_0001, Odoo aggiunge __import__.). In alternativa usare la colonna partner_id con il NOME esatto. Per l\'azienda: base.main_company -> usare base.main_partner.', 20, None),
            ('account_number', 'Numero conto (IBAN)', 'O', 'Testo', 'IT60X0542811101000000123456', 'IBAN senza spazi. Odoo ne verifica il formato e ricava il paese.', 32, None),
            ('holder_name', 'Intestatario (se diverso)', 'F', 'Testo', '', 'Solo se diverso dal nome del titolare.', 22, None),
            ('bank_name', 'Nome banca', 'C', 'Testo', 'Intesa Sanpaolo', 'In Odoo 20 la banca è un testo sul conto (il modello res.bank non esiste più).', 22, None),
            ('bank_bic', 'BIC/SWIFT', 'C', 'Testo 8/11', 'BCITITMM', 'Usato nei bonifici SEPA/esteri.', 12, None),
            ('allow_out_payment', 'Pagamenti in uscita consentiti', 'C', 'Booleano', 'TRUE / FALSE', 'Se FALSE Odoo blocca i bonifici SEPA verso questo IBAN (conto "non fidato"). Richiede il gruppo "Validare conti bancari".', 14, BOOL),
            ('note', 'Note', 'F', 'Testo', '', '', 20, None),
        ],
        'examples': [
            ['IBAN_FORN_0001', 'FORN_0001', 'IT60X0542811101000000123456', '', 'ESEMPIO Intesa Sanpaolo', 'BCITITMM', 'TRUE', ''],
            ['IBAN_FORN_0002', 'FORN_0002', 'IT02L1234512345123456789012', '', 'ESEMPIO UniCredit', 'UNCRITMM', 'TRUE', ''],
        ],
    },
    {
        'file': '03_Dipendenti.xlsx',
        'model': 'hr.employee',
        'title': 'Dipendenti',
        'menu': 'Dipendenti > Dipendenti',
        'cols': [
            ('id', 'ID esterno', 'C', 'Testo', 'DIP_0001', EXT_ID, 12, None),
            ('name', 'Nome dipendente', 'O', 'Testo', 'Cognome Nome', 'Usare sempre lo stesso ordine (consiglio: Cognome Nome).', 24, None),
            ('company_id', 'Azienda', 'O', 'Nome', 'Ragione sociale azienda', 'Nome ESATTO dell\'azienda in Odoo (multi-company). Vuoto = azienda corrente.', 20, None),
            ('work_email', 'Email lavoro', 'C', 'Testo', 'm.rossi@azienda.it', '', 24, None),
            ('work_phone', 'Telefono lavoro', 'F', 'Testo', '', '', 14, None),
            ('mobile_phone', 'Cellulare lavoro', 'F', 'Testo', '', '', 14, None),
            ('department_id', 'Reparto', 'C', 'Nome', 'Amministrazione', 'Nome esatto del reparto (crearli prima, o confermare la creazione in import).', 18, None),
            ('job_id', 'Posizione lavorativa', 'F', 'Nome', 'Contabile', 'Anagrafica hr.job.', 16, None),
            ('job_title', 'Titolo', 'F', 'Testo', 'Responsabile contabilità', 'Testo libero mostrato sulla scheda.', 22, None),
            ('parent_id/id', 'Responsabile (ID esterno)', 'F', 'ID esterno', 'DIP_0001', 'ID esterno di un altro dipendente. Il responsabile deve trovarsi in una riga PRECEDENTE o essere già in Odoo.', 14, None),
            ('coach_id/id', 'Coach (ID esterno)', 'F', 'ID esterno', '', 'Come Responsabile.', 12, None),
            ('resource_calendar_id', 'Orario di lavoro', 'C', 'Nome', 'Standard 40 hours/week', 'Nome esatto del calendario (Configurazione > Orari di lavoro).', 20, None),
            ('work_location_id', 'Luogo di lavoro', 'F', 'Nome', 'Sede', '', 12, None),
            ('employee_type_id/id', 'Tipo dipendente', 'C', 'ID esterno', 'hr.contract_type_employee', 'Odoo 20: anagrafica hr.employee.type. Vedi foglio Valori.', 26, [e[0] for e in EMP_TYPES]),
            ('contract_date_start', 'Data inizio contratto', 'C', 'Data', '2024-01-15', 'Formato AAAA-MM-GG.', 14, None),
            ('date_version', 'Data validità dati', 'F', 'Data', '2024-01-15', 'Odoo 20: data di efficacia della "versione" del dipendente. Vuoto = oggi.', 14, None),
            ('identification_id', 'N. identificazione (Codice Fiscale)', 'C', 'Testo', 'RSSMRA80A01H501U', 'Convenzione consigliata: Codice Fiscale.', 18, None),
            ('sex', 'Sesso (legale)', 'F', 'Selezione', 'male / female', 'Sesso legale anagrafico (Odoo 20 ammette solo male/female).', 10, ['male', 'female']),
            ('birthday', 'Data di nascita', 'F', 'Data', '1980-01-01', '', 12, None),
            ('place_of_birth', 'Luogo di nascita', 'F', 'Testo', 'Roma', '', 14, None),
            ('country_of_birth/id', 'Paese di nascita', 'F', 'ID esterno', 'base.it', '', 10, None),
            ('country_id/id', 'Nazionalità', 'F', 'ID esterno', 'base.it', '', 10, None),
            ('private_street', 'Indirizzo privato', 'F', 'Testo', '', '', 22, None),
            ('private_zip', 'CAP privato', 'F', 'Testo', '', '', 8, None),
            ('private_city', 'Città privata', 'F', 'Testo', '', '', 14, None),
            ('private_state_id/id', 'Provincia privata', 'F', 'ID esterno', 'base.state_it_rm', PROVINCE, 14, None),
            ('private_country_id/id', 'Paese privato', 'F', 'ID esterno', 'base.it', '', 10, None),
            ('private_email', 'Email privata', 'F', 'Testo', '', '', 20, None),
            ('private_phone', 'Telefono privato', 'F', 'Testo', '', '', 14, None),
            ('marital', 'Stato civile', 'F', 'Selezione', 'single', 'Default single.', 12, ['single', 'married', 'cohabitant', 'widower', 'divorced']),
            ('children', 'Figli a carico', 'F', 'Intero', '0', '', 8, None),
            ('emergency_contact', 'Contatto emergenza', 'F', 'Testo', '', '', 18, None),
            ('emergency_phone', 'Tel. emergenza', 'F', 'Testo', '', '', 14, None),
            ('category_ids', 'Etichette', 'F', 'Nomi separati da virgola', '', '', 16, None),
            ('barcode', 'Badge ID', 'F', 'Testo', '', 'Univoco.', 10, None),
        ],
        'examples': [
            ['DIP_0001', 'ESEMPIO Bianchi Laura', '', 'l.bianchi@example.com', '', '', 'Amministrazione', '', 'Responsabile amministrativa', '', '', 'Standard 40 hours/week', '', 'hr.contract_type_employee', '2020-03-01', '', 'BNCLRA85C41F205X', 'female', '1985-03-01', 'Milano', 'base.it', 'base.it', '', '', '', '', 'base.it', '', '', 'married', 1, '', '', '', ''],
            ['DIP_0002', 'ESEMPIO Verdi Paolo', '', 'p.verdi@example.com', '', '', 'Amministrazione', '', 'Contabile', 'DIP_0001', '', 'Standard 40 hours/week', '', 'hr.contract_type_employee', '2023-09-01', '', 'VRDPLA90A01L219Z', 'male', '1990-01-01', 'Torino', 'base.it', 'base.it', '', '', '', '', 'base.it', '', '', 'single', 0, '', '', '', ''],
        ],
    },
    {
        'file': '04_Piano_dei_Conti.xlsx',
        'model': 'account.account',
        'title': 'Piano dei conti',
        'menu': 'Contabilità > Configurazione > Piano dei conti',
        'cols': [
            ('id', 'ID esterno', 'C', 'Testo', 'CONTO_6105', 'Per AGGIORNARE conti già esistenti (es. quelli di l10n_it) esportarli prima con "Compatibile con importazione" e usare l\'ID esportato.', 14, None),
            ('code', 'Codice', 'O', 'Testo', '6105', 'Univoco per azienda. Solo cifre/lettere/punti. Formattare la colonna come TESTO.', 10, None),
            ('name', 'Nome conto', 'O', 'Testo', 'Consulenze amministrative', '', 34, None),
            ('account_type', 'Tipo', 'O', 'Selezione', 'expense', 'Vedi foglio Valori. Determina bilancio/CE, chiusura e riporto saldi.', 22, [a[0] for a in ACCOUNT_TYPES]),
            ('reconcile', 'Consenti riconciliazione', 'C', 'Booleano', 'TRUE / FALSE', 'OBBLIGATORIO TRUE per asset_receivable e liability_payable. Consigliato TRUE per conti transitori/anticipi.', 12, BOOL),
            ('non_trade', 'Non commerciale', 'F', 'Booleano', 'FALSE', 'Solo per crediti/debiti non commerciali (es. IVA, erario) di tipo receivable/payable.', 10, BOOL),
            ('tag_ids/id', 'Etichette (riclassifica CEE)', 'C', 'ID esterni separati da virgola', 'l10n_it.account_tag_B_PL', 'Riclassifica bilancio CEE della localizzazione italiana. Vedi foglio Valori.', 26, [t[0] for t in IT_TAGS]),
            ('currency_id', 'Valuta conto', 'F', 'Codice ISO', '', 'Solo per conti in valuta estera (es. banca USD).', 8, None),
            ('description', 'Descrizione', 'F', 'Testo', '', '', 20, None),
            ('note', 'Note interne', 'F', 'Testo', '', '', 20, None),
        ],
        'examples': [
            ['CONTO_6105', '6105', 'ESEMPIO Consulenze amministrative', 'expense', 'FALSE', 'FALSE', 'l10n_it.account_tag_B_PL', '', '', ''],
            ['CONTO_1520', '1520', 'ESEMPIO Anticipi a fornitori', 'asset_current', 'TRUE', 'FALSE', 'l10n_it.account_tag_C_ATT', '', '', ''],
            ['CONTO_4105', '4105', 'ESEMPIO Ricavi servizi', 'income', 'FALSE', 'FALSE', 'l10n_it.account_tag_A_PL', '', '', ''],
        ],
    },
    {
        'file': '05_Budget_Analitico.xlsx',
        'model': 'budget.analytic',
        'title': 'Budget analitico (testata + righe)',
        'menu': 'Contabilità > Contabilità > Budget  (modulo Enterprise account_budget)',
        'one2many': True,
        'cols': [
            ('id', 'ID esterno', 'C', 'Testo', 'BDG_2026_COMM', 'Solo sulla PRIMA riga di ogni budget. Righe successive: vuoto.', 16, None),
            ('name', 'Nome budget', 'O', 'Testo', 'Budget 2026 Commesse', 'Solo sulla PRIMA riga di ogni budget.', 24, None),
            ('date_from', 'Data inizio', 'O', 'Data', '2026-01-01', 'Periodo dell\'INTERO budget: le righe ereditano queste date (non hanno date proprie). Per budget mensili creare un budget per mese.', 12, None),
            ('date_to', 'Data fine', 'O', 'Data', '2026-12-31', '', 12, None),
            ('budget_type', 'Tipo budget', 'O', 'Selezione', 'expense', 'expense = costi, revenue = ricavi, both = entrambi (in both: ricavi +, costi -).', 10, ['expense', 'revenue', 'both']),
            ('user_id', 'Responsabile', 'F', 'Nome utente', '', 'Nome esatto dell\'utente Odoo.', 16, None),
            ('company_id', 'Azienda', 'F', 'Nome', '', 'Vuoto = azienda corrente.', 16, None),
            ('budget_line_ids/account_id', 'Riga: conto analitico piano principale', 'C', 'Nome o codice conto analitico', 'Commessa ALFA', 'Colonna del PRIMO piano analitico (di default "Progetto"). Almeno un conto analitico per riga.', 26, None),
            ('budget_line_ids/x_plan2_id', 'Riga: conto analitico piano 2', 'F', 'Nome o codice conto analitico', 'Produzione', 'I piani aggiuntivi hanno colonne x_planN_id dove N = ID del piano: VERIFICARE sull\'istanza (vedi manuale) e rinominare/duplicare la colonna.', 22, None),
            ('budget_line_ids/budget_amount', 'Riga: importo a budget', 'O', 'Numero', '50000', 'Importo positivo per budget expense/revenue.', 14, None),
        ],
        'examples': [
            ['BDG_2026_COMM', 'ESEMPIO Budget 2026 Commesse', '2026-01-01', '2026-12-31', 'expense', '', '', 'Commessa ALFA', '', 50000],
            ['', '', '', '', '', '', '', 'Commessa BETA', '', 32000],
            ['', '', '', '', '', '', '', 'Commessa GAMMA', 'Produzione', 18000],
            ['BDG_2026_Q1_CDC', 'ESEMPIO Budget Q1 2026 Centri di costo', '2026-01-01', '2026-03-31', 'expense', '', '', '', 'Produzione', 120000],
            ['', '', '', '', '', '', '', '', 'Amministrazione', 45000],
        ],
    },
    {
        'file': '06_Budget_Finanziario.xlsx',
        'model': 'account.report.budget',
        'title': 'Budget finanziario (per conto contabile, nei report CE/SP)',
        'menu': 'Contabilità > Reportistica > Conto economico > (pulsante Budget)  -  modello account.report.budget',
        'one2many': True,
        'cols': [
            ('id', 'ID esterno', 'C', 'Testo', 'BUDFIN_2026', 'Solo sulla PRIMA riga del budget.', 14, None),
            ('name', 'Nome budget', 'O', 'Testo', 'Budget 2026', 'Solo sulla PRIMA riga. È il nome che si seleziona nel report.', 18, None),
            ('company_id', 'Azienda', 'O', 'Nome', '', 'Solo prima riga. Vuoto = azienda corrente.', 16, None),
            ('item_ids/account_id', 'Riga: conto contabile', 'O', 'Codice conto', '4101', 'Codice del conto del piano dei conti (deve esistere).', 12, None),
            ('item_ids/date', 'Riga: data (mese)', 'O', 'Data', '2026-01-01', 'Una riga per conto per mese. Convenzione: primo giorno del mese.', 12, None),
            ('item_ids/amount', 'Riga: importo', 'O', 'Numero', '-10000', 'Convenzione SALDO CONTABILE (dare - avere): COSTI positivi, RICAVI negativi. Verificare con un test (vedi manuale).', 12, None),
        ],
        'examples': [
            ['BUDFIN_2026', 'ESEMPIO Budget 2026', '', '4101', '2026-01-01', -10000],
            ['', '', '', '4101', '2026-02-01', -12000],
            ['', '', '', '6105', '2026-01-01', 1500],
            ['', '', '', '6105', '2026-02-01', 1500],
        ],
    },
]


def style_header(ws, cols):
    fills = {'O': RED, 'C': BLUE, 'F': GREY}
    for i, (field, label, lvl, typ, fmt, note, width, _vals) in enumerate(cols, 1):
        c = ws.cell(row=1, column=i, value=field)
        c.fill = fills[lvl]
        c.font = HEAD_FONT
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        c.border = THIN
        c.comment = Comment(f'{label}\n[{ {"O": "OBBLIGATORIO", "C": "Consigliato", "F": "Facoltativo"}[lvl]}]\n'
                            f'Formato: {fmt}\n{note}', 'Template', width=320, height=160)
        ws.column_dimensions[get_column_letter(i)].width = max(width, len(field) + 2)
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = 'A2'


def add_validations(ws, cols, max_row=2000):
    long_lists = [c[0] for c in cols if c[7] and len(','.join(c[7])) >= 250]  # stesso ordine di values_sheet
    for i, col in enumerate(cols, 1):
        vals = col[7]
        if not vals:
            continue
        letter = get_column_letter(i)
        if len(','.join(vals)) < 250:
            dv = DataValidation(type='list', formula1='"%s"' % ','.join(vals), allow_blank=True)
        else:  # lista lunga -> riferimento al foglio Valori
            vcol = get_column_letter(1 + 3 * long_lists.index(col[0]))
            dv = DataValidation(type='list', formula1=f"=Valori!${vcol}$2:${vcol}${len(vals) + 1}", allow_blank=True)
        dv.showErrorMessage = False  # avviso morbido: consente valori custom (es. tag multipli)
        ws.add_data_validation(dv)
        dv.add(f'{letter}2:{letter}{max_row}')


def legend_sheet(wb, tpl):
    ws = wb.create_sheet('Legenda')
    ws.append([f'{tpl["title"]}  -  modello Odoo: {tpl["model"]}'])
    ws['A1'].font = Font(bold=True, size=13)
    ws.append([f'Menu di importazione: {tpl["menu"]}  >  Preferiti (stella) / ingranaggio > Importa record'])
    ws.append(['Rosso = obbligatorio · Blu = consigliato · Grigio = facoltativo. Le righe gialle nel foglio Dati sono ESEMPI: cancellarle prima di importare.'])
    if tpl.get('one2many'):
        ws.append(['STRUTTURA TESTATA/RIGHE: la testata (id, nome, date...) si compila SOLO sulla prima riga; le righe successive hanno le colonne di testata VUOTE e compilano solo le colonne con "/" (righe di dettaglio).'])
    ws.append([])
    head = ['Colonna (nome tecnico)', 'Significato', 'Obbligatorietà', 'Tipo', 'Formato / esempio', 'Note']
    ws.append(head)
    hr = ws.max_row
    for c in ws[hr]:
        c.fill = BLUE
        c.font = HEAD_FONT
    lvl = {'O': 'OBBLIGATORIO', 'C': 'Consigliato', 'F': 'Facoltativo'}
    for field, label, l, typ, fmt, note, _w, _v in tpl['cols']:
        ws.append([field, label, lvl[l], typ, fmt, note])
        if l == 'O':
            ws.cell(row=ws.max_row, column=3).font = Font(bold=True, color='C0392B')
    for col, w in zip('ABCDEF', (34, 32, 14, 16, 30, 80)):
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=hr + 1):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical='top')
            c.border = THIN


def values_sheet(wb, tpl):
    lists = []
    for field, label, *_rest, vals in tpl['cols']:
        if vals and field == 'account_type':
            lists.append((field, ACCOUNT_TYPES))
        elif vals and field == 'tag_ids/id':
            lists.append((field, IT_TAGS))
        elif vals and field == 'employee_type_id/id':
            lists.append((field, EMP_TYPES))
        elif vals and 'payment_term' in field:
            lists.append((field, PAYMENT_TERMS))
    if not lists:
        return
    ws = wb.create_sheet('Valori')
    col = 1
    for field, pairs in lists:
        ws.cell(row=1, column=col, value=field).font = Font(bold=True)
        ws.cell(row=1, column=col + 1, value='Descrizione').font = Font(bold=True)
        for r, (k, d) in enumerate(pairs, 2):
            ws.cell(row=r, column=col, value=k)
            ws.cell(row=r, column=col + 1, value=d)
        ws.column_dimensions[get_column_letter(col)].width = 44
        ws.column_dimensions[get_column_letter(col + 1)].width = 46
        col += 3


def coa_reference_sheet(wb):
    if not os.path.exists(L10N_IT_COA):
        return
    ws = wb.create_sheet('Rif_PdC_l10n_it')
    ws.append(['Piano dei conti standard Odoo l10n_it (già creato all\'installazione della localizzazione). '
               'Usarlo per mappare i conti del vecchio gestionale: importare SOLO i conti mancanti.'])
    ws['A1'].font = Font(bold=True)
    ws.append(['code', 'Nome IT', 'account_type', 'reconcile', 'non_trade', 'tag_ids/id'])
    for c in ws[2]:
        c.font = HEAD_FONT
        c.fill = BLUE
    with open(L10N_IT_COA, encoding='utf-8') as f:
        for r in csv.DictReader(f):
            ws.append([r['code'], r['name@it'], r['account_type'], r['reconcile'], r['non_trade'], r['tag_ids']])
    for col, w in zip('ABCDEF', (10, 50, 20, 10, 10, 30)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = 'A3'


def matrix_sheet(wb):
    """Foglio di lavoro (NON importabile) per preparare il budget finanziario in formato conto x mese."""
    ws = wb.create_sheet('Matrice_preparazione')
    ws.append(['Foglio di PREPARAZIONE (non si importa). Compilare conto x mese, poi convertire nel foglio Dati con lo script '
               'budget_matrice_to_import.py oppure copiando a mano (una riga per conto e mese).'])
    ws['A1'].font = Font(bold=True)
    months = [f'2026-{m:02d}-01' for m in range(1, 13)]
    ws.append(['code', 'Nome conto (promemoria)'] + months)
    for c in ws[2]:
        c.font = HEAD_FONT
        c.fill = BLUE
    ws.append(['4101', 'ESEMPIO Ricavi vendite'] + [-10000] * 12)
    ws.append(['6105', 'ESEMPIO Consulenze'] + [1500] * 12)
    ws.column_dimensions['B'].width = 30
    for i in range(3, 15):
        ws.column_dimensions[get_column_letter(i)].width = 12
    ws.freeze_panes = 'C3'


def build(tpl):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Dati'  # PRIMO foglio = quello proposto da Odoo all'import
    cols = tpl['cols']
    style_header(ws, cols)
    text_cols = {i for i, c in enumerate(cols, 1) if c[0] in ('code', 'zip', 'vat', 'account_number', 'l10n_it_pa_index', 'private_zip', 'item_ids/account_id', 'ref', 'barcode')}
    for row in tpl['examples']:
        ws.append(row)
        for i, c in enumerate(ws[ws.max_row], 1):
            c.fill = EXAMPLE
            c.border = THIN
            if i in text_cols:
                c.number_format = '@'
    for i in text_cols:  # colonne testo anche per le righe vuote
        for r in range(ws.max_row + 1, 501):
            ws.cell(row=r, column=i).number_format = '@'
    add_validations(ws, cols)
    legend_sheet(wb, tpl)
    values_sheet(wb, tpl)
    if tpl['model'] == 'account.account':
        coa_reference_sheet(wb)
    if tpl['model'] == 'account.report.budget':
        matrix_sheet(wb)
    path = os.path.join(HERE, tpl['file'])
    wb.save(path)
    return path


if __name__ == '__main__':
    for t in TEMPLATES:
        print('creato', build(t))
