#!/usr/bin/env python3
"""Genera il manuale DOCX dei template di importazione Odoo."""
import os

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from genera_template import TEMPLATES

HERE = os.path.dirname(os.path.abspath(__file__))
doc = Document()
st = doc.styles['Normal']
st.font.name = 'Calibri'
st.font.size = Pt(10.5)


def shade(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    tcPr.append(shd)


def table(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Table Grid'
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        c.text = ''
        r = c.paragraphs[0].add_run(h)
        r.bold = True
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        shade(c, '2E5C8A')
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = str(v)
    for row in t.rows:
        for c in row.cells:
            for p in c.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(9)
    doc.add_paragraph()
    return t


def p(text, bold_prefix=None):
    para = doc.add_paragraph()
    if bold_prefix:
        para.add_run(bold_prefix).bold = True
    para.add_run(text)
    return para


def bullet(text):
    doc.add_paragraph(text, style='List Bullet')


def num(text):
    doc.add_paragraph(text, style='List Number')


# ---------------------------------------------------------------- copertina
h = doc.add_heading('Manuale template di importazione dati Odoo', 0)
p('Fornitori · Banche · Dipendenti · Piano dei conti · Budget analitico · Budget finanziario')
p('Versione di riferimento: Odoo 20.0 Enterprise (repository pixora/odoo, branch 20.0). Differenze con 18/19 al capitolo 9. '
  'Classificazione: Interno.')
p('I nomi dei campi Community (contatti, conti bancari, dipendenti, piano dei conti) sono verificati sul sorgente '
  'Odoo 20.0 del repository. Il budget analitico (budget.analytic, Enterprise) è confermato su Odoo 20 dal modulo '
  'demo pix_demo_family_office (repo pixora/demo-alina); il budget finanziario (account.report.budget) è verificato '
  'su Odoo 19 e va riconfermato su 20 con un import di prova. Prima del go-live eseguire sempre un '
  'import di prova in un database di staging/test.')

# ---------------------------------------------------------------- 1
doc.add_heading('1. Contenuto del pacchetto', 1)
table(['File', 'Modello Odoo', 'Cosa carica'], [
    [t['file'], t['model'], t['title']] for t in TEMPLATES
] + [
    ['budget_matrice_to_import.py', '-', 'Converte la matrice conto x mese del budget finanziario nel formato importabile'],
    ['genera_template.py / genera_manuale.py', '-', 'Rigenerano template e manuale (per adattarli a un cliente)'],
])
p('Ogni file Excel contiene:')
bullet('Foglio "Dati" (PRIMO foglio): è quello che si importa. Riga 1 = nomi tecnici dei campi. '
       'Le righe gialle sono ESEMPI fittizi da cancellare prima dell\'import.')
bullet('Foglio "Legenda": significato di ogni colonna, obbligatorietà, formato e note.')
bullet('Foglio "Valori" (dove serve): codici ammessi per i campi a scelta (tipi conto, etichette CEE, termini di pagamento).')
bullet('Colori intestazione: ROSSO = obbligatorio, BLU = consigliato, GRIGIO = facoltativo. '
       'Passando il mouse sull\'intestazione compare un commento con la spiegazione.')

# ---------------------------------------------------------------- 2
doc.add_heading('2. Regole generali dell\'import Odoo', 1)
table(['Regola', 'Come applicarla'], [
    ['Nomi colonna tecnici', 'Non rinominare le intestazioni: Odoo le riconosce automaticamente in qualsiasi lingua. '
                             'Le colonne non necessarie si possono eliminare (tranne quelle rosse).'],
    ['ID esterno (colonna id)', 'Codice univoco scelto da voi (FORN_0001, DIP_0001...). Odoo lo salva come __import__.FORN_0001. '
                                'Reimportando un file con lo stesso id il record viene AGGIORNATO, non duplicato.'],
    ['Collegamenti con /id', 'Colonne che terminano in /id (es. partner_id/id) vogliono l\'ID esterno del record collegato: '
                             'quelli creati da voi (FORN_0001) o quelli standard (base.it, base.state_it_mi).'],
    ['Collegamenti per nome', 'Colonne senza /id (es. department_id, company_id) vogliono il NOME esatto come appare in Odoo '
                              '(o il codice per i conti contabili). Nomi ambigui o con refusi = errore.'],
    ['Date', 'Formato AAAA-MM-GG (2026-01-31) oppure cella data di Excel.'],
    ['Booleani', 'TRUE / FALSE (accettati anche 1/0, Vero/Falso).'],
    ['Numeri', 'Senza separatore delle migliaia, decimali con il punto o cella numerica Excel.'],
    ['Colonne testo', 'CAP, codici conto, P.IVA, IBAN: la colonna è già formattata come Testo per non perdere gli zeri iniziali.'],
    ['Più valori', 'Per campi multipli (etichette) separare con la virgola: Materie prime,Strategico'],
    ['Testata + righe', 'Nei budget: compilare la testata solo nella prima riga, le righe successive solo nelle colonne con "/" (vedi cap. 6-7).'],
    ['Formule', 'Evitarle nel foglio Dati: incollare come valori prima dell\'import.'],
])

# ---------------------------------------------------------------- 3
doc.add_heading('3. Ordine di caricamento consigliato', 1)
p('Alcuni dati dipendono da altri: rispettare questa sequenza evita errori di "record non trovato".')
table(['#', 'File', 'Prerequisiti'], [
    ['0', '(configurazione)', 'Azienda/e, localizzazione italiana (l10n_it + l10n_it_edi) installata, valute e lingue attive, '
                              'posizioni fiscali, reparti e orari di lavoro, piani e conti analitici.'],
    ['1', '04_Piano_dei_Conti', 'Solo i conti MANCANTI rispetto al piano standard l10n_it (vedi foglio Rif_PdC_l10n_it).'],
    ['2', '01_Fornitori', 'Conto debiti, termini di pagamento e posizioni fiscali esistenti.'],
    ['3', '02_Conti_Bancari', 'Fornitori (2) già importati.'],
    ['4-5', '03_Dipendenti', 'Reparti, posizioni, orari di lavoro. IBAN dipendenti: vedi 5.2.'],
    ['6', '05_Budget_Analitico', 'Modulo Budget (account_budget) attivo, piani e conti analitici esistenti.'],
    ['7', '06_Budget_Finanziario', 'Piano dei conti definitivo.'],
])

# ---------------------------------------------------------------- 4
doc.add_heading('4. Procedura di import passo-passo', 1)
for s in [
    'Compilare il foglio Dati e cancellare le righe gialle di esempio.',
    'Aprire il menu indicato nel foglio Legenda (es. Contabilità > Fornitori > Fornitori), vista elenco.',
    'Cliccare l\'ingranaggio/icona azioni accanto al titolo > Importa record > Carica file.',
    'Se il file ha più fogli, verificare che sia selezionato il foglio "Dati".',
    'Controllare che ogni colonna sia abbinata al campo corretto (le colonne tecniche si abbinano da sole).',
    'Cliccare TESTA: Odoo simula l\'import e mostra gli errori riga per riga senza scrivere nulla.',
    'Correggere il file e ripetere il test finché non ci sono errori, poi cliccare IMPORTA.',
    'Verificare a campione 3-5 record importati e conservare il file usato (evidenza di migrazione).',
]:
    num(s)
p('Suggerimento: per file grandi (> 2.000 righe) usare l\'opzione "Importa in lotti" nelle impostazioni avanzate dell\'import.')

# ---------------------------------------------------------------- 5 anagrafiche
doc.add_heading('5. Anagrafiche', 1)
doc.add_heading('5.1 Fornitori (01_Fornitori.xlsx - res.partner)', 2)
p('Una riga per fornitore. Campi chiave per la fatturazione elettronica italiana: P.IVA (vat) con prefisso IT, '
  'Codice Fiscale, indirizzo completo (via, CAP, città, provincia, nazione). PEC e Codice destinatario servono '
  'soprattutto se si emettono documenti verso il fornitore.')
bullet('supplier_rank = 1 è indispensabile: senza, il contatto non compare nel menu Fornitori.')
bullet('Provincia: ID esterno base.state_it_ + sigla in minuscolo (Milano = base.state_it_mi).')
bullet('Termini di pagamento: usare gli ID standard del foglio Valori oppure, per termini personalizzati, '
       'rinominare la colonna in property_supplier_payment_term_id e scrivere il nome esatto.')
bullet('La P.IVA viene validata: un checksum errato blocca la riga. Correggere il dato alla fonte, non disattivare il controllo.')
bullet('Lo stesso file, aggiungendo customer_rank = 1, può caricare anche i clienti.')

doc.add_heading('5.2 Banche (02_Conti_Bancari.xlsx - res.partner.bank)', 2)
p('In Odoo 20 non esiste più un\'anagrafica separata degli istituti bancari (modello res.bank): nome banca e BIC '
  'sono campi di testo direttamente sul conto bancario. Un unico file quindi carica IBAN + banca, collegati al titolare '
  '(fornitore, dipendente o azienda). Il paese del conto viene ricavato automaticamente dall\'IBAN.')
p('', 'Conti bancari dei DIPENDENTI:')
num('Importare i dipendenti (03_Dipendenti).')
num('Importare gli IBAN in 02_Conti_Bancari usando la colonna partner_id (senza /id) con il NOME del dipendente '
    '(Odoo crea per ogni dipendente un contatto con lo stesso nome).')
num('Reimportare 03_Dipendenti con le sole colonne id e bank_account_ids contenente l\'IBAN.')
p('', 'Conti bancari dell\'AZIENDA:')
p('Per i conti dell\'azienda è preferibile creare direttamente i registri banca (Contabilità > Configurazione > Registri, '
  'tipo Banca, campo Numero conto): Odoo crea il conto contabile (prefisso 182) e il conto bancario collegato.')
p('allow_out_payment (Pagamenti in uscita consentiti): se FALSE Odoo blocca i bonifici verso quell\'IBAN. '
  'Impostarlo a TRUE solo per IBAN verificati (presidio anti-frode, coerente con i controlli ISO 27001).')

doc.add_heading('5.3 Dipendenti (03_Dipendenti.xlsx - hr.employee)', 2)
p('Una riga per dipendente. Il Responsabile (parent_id/id) deve essere in una riga precedente o già presente in Odoo: '
  'ordinare il file dall\'alto (direzione) verso il basso.')
bullet('In Odoo 20 i dati contrattuali/anagrafici sono "versionati" (hr.version): date_version indica da quando valgono. '
       'In import si può lasciare vuoto (= oggi) o mettere la data di assunzione.')
bullet('company_id: nome esatto dell\'azienda (importante in multi-company).')
bullet('Codice fiscale: convenzione consigliata nel campo identification_id (il campo ssnid non esiste più in Odoo 20).')
bullet('Tipo dipendente: in Odoo 20 è un\'anagrafica (employee_type_id) e non più un elenco fisso: usare gli ID del foglio Valori (es. hr.contract_type_employee).')
p('', 'Protezione dati (GDPR):')
p('Il file contiene dati personali. Applicare la minimizzazione: compilare solo le colonne realmente necessarie '
  '(le facoltative si possono eliminare). Non inserire dati sanitari o categorie particolari. Trasmettere il file '
  'solo tramite canali aziendali approvati, non via email non cifrata, e cancellarlo al termine della migrazione '
  'secondo la policy di retention. Per trattamenti nuovi o estesi coinvolgere il DPO.')

doc.add_heading('5.4 Piano dei conti (04_Piano_dei_Conti.xlsx - account.account)', 2)
p('Installando la localizzazione italiana Odoo crea già circa 190 conti (elencati nel foglio Rif_PdC_l10n_it). '
  'Il metodo consigliato è MAPPARE i conti del vecchio gestionale su quelli standard e importare solo quelli mancanti.')
table(['Campo', 'Regola'], [
    ['code', 'Univoco per azienda. Lunghezza coerente con lo standard (4 cifre in l10n_it, es. 4204). I conti banca usano prefisso 182, cassa 180.'],
    ['account_type', 'Determina se il conto va in Stato patrimoniale o Conto economico, la chiusura d\'esercizio e il riporto saldi. Vedi foglio Valori.'],
    ['reconcile', 'TRUE obbligatorio per asset_receivable e liability_payable; consigliato per conti transitori, anticipi, effetti.'],
    ['tag_ids/id', 'Riclassifica del bilancio CEE (l10n_it.account_tag_...). Fondamentale per i report legali italiani.'],
    ['equity_unaffected', 'Uno solo per azienda: esiste già, non importarne un altro.'],
])
p('Per modificare conti già esistenti: esportarli (Esporta > "Compatibile con importazione", includendo id), modificare e reimportare.')

# ---------------------------------------------------------------- 6 budget analitico
doc.add_heading('6. Budget analitico (05_Budget_Analitico.xlsx - budget.analytic)', 1)
p('Il budget analitico confronta importi previsti e consuntivi per CONTO ANALITICO (commesse, centri di costo, progetti). '
  'Richiede il modulo Enterprise Budget (account_budget).')
p('', 'Struttura testata + righe:')
table(['Riga Excel', 'id / name / date / budget_type', 'budget_line_ids/...'], [
    ['2 (prima riga budget A)', 'COMPILATE', 'prima riga di dettaglio'],
    ['3', 'vuote', 'seconda riga di dettaglio di A'],
    ['4', 'vuote', 'terza riga di dettaglio di A'],
    ['5 (prima riga budget B)', 'COMPILATE', 'prima riga di dettaglio di B'],
])
bullet('Le righe NON hanno date proprie: ereditano date_from/date_to della testata. Per un budget mensile o trimestrale '
       'creare un budget per periodo (es. "Budget 2026-01", "Budget 2026-02"...).')
bullet('budget_type: expense (costi), revenue (ricavi), both (entrambi: ricavi positivi, costi negativi). '
       'Per expense e revenue inserire importi positivi.')
bullet('Il budget viene importato in stato Bozza: confermarlo (pulsante Apri) per vedere il consuntivo.')
p('', 'Colonne dei piani analitici:')
p('Ogni piano analitico ha una sua colonna. Il primo piano (di norma "Progetto") usa budget_line_ids/account_id; '
  'gli altri piani usano budget_line_ids/x_planN_id dove N è l\'ID del piano. Il nome cambia da database a database. '
  'Per scoprirlo: Contabilità > Configurazione > Piani analitici, aprire il piano e leggere l\'ID nell\'URL '
  '(es. .../analytic.plan/2 -> x_plan2_id); oppure esportare una riga di budget e leggere i nomi colonna. '
  'Aggiungere/rinominare le colonne del template di conseguenza. Il valore è il nome (o codice) del conto analitico.')

# ---------------------------------------------------------------- 7 budget finanziario
doc.add_heading('7. Budget finanziario (06_Budget_Finanziario.xlsx - account.report.budget)', 1)
p('Il budget finanziario è definito per CONTO CONTABILE e per MESE e si visualizza direttamente nei report '
  'Conto economico / Stato patrimoniale (pulsante Budget) a confronto con il consuntivo.')
bullet('Testata: id, name, company_id solo nella prima riga.')
bullet('Righe: una per conto per mese (item_ids/account_id = codice conto, item_ids/date = primo giorno del mese, item_ids/amount).')
bullet('Segno degli importi: le righe seguono la convenzione del saldo contabile (dare - avere): COSTI positivi, RICAVI negativi. '
       'Prima del caricamento massivo fare una prova: inserire a mano un valore nel report (modalità budget), '
       'esportare il budget e verificare il segno salvato; allineare il file di conseguenza.')
p('', 'Preparazione comoda in formato matrice:')
p('Il foglio Matrice_preparazione permette di compilare il budget come una tabella conto x 12 mesi. '
  'Lo script budget_matrice_to_import.py lo converte nel formato riga-per-mese:')
code = doc.add_paragraph('python3 budget_matrice_to_import.py 06_Budget_Finanziario.xlsx "Budget 2026" BUDFIN_2026')
code.runs[0].font.name = 'Consolas'
p('Il risultato (06_Budget_Finanziario_IMPORT.xlsx) si importa direttamente.')

# ---------------------------------------------------------------- 8 alternative
doc.add_heading('8. Modalità di caricamento: tre alternative', 1)
table(['Opzione', 'Come', 'Pro', 'Contro', 'Quando'], [
    ['A. Import da interfaccia (consigliata)', 'Questi template + Importa record', 'Nessuno sviluppo, test integrato, usabile dal key user',
     'Manuale, file grandi lenti, ordine dipendenze a carico dell\'utente', 'Avvio progetto, volumi < 10.000 righe per file'],
    ['B. Import con ID esterni ripetibile', 'Stessi template, ma id sempre compilato e reimport iterativi', 'Migrazioni di prova ripetibili, aggiornamenti senza duplicati',
     'Disciplina sui codici id, conservare i file', 'Più cicli di migrazione (test -> UAT -> produzione)'],
    ['C. Script via API (XML-RPC / JSON-2)', 'Script Python che legge gli Excel e chiama create/write', 'Automatizzabile, validazioni custom, grandi volumi',
     'Richiede sviluppo e test, gestione chiavi API e log', 'Volumi alti, migrazioni ricorrenti o da più fonti'],
])

# ---------------------------------------------------------------- 9 versioni
doc.add_heading('9. Differenze tra versioni (18 / 19 / 20)', 1)
table(['Area', 'Odoo 18', 'Odoo 19', 'Odoo 20 (questi template)'], [
    ['Anagrafica banche', 'res.bank + bank_id', 'res.bank + bank_id', 'non esiste: bank_name, bank_bic sul conto'],
    ['Numero conto (IBAN)', 'acc_number', 'acc_number', 'account_number'],
    ['Intestatario conto', 'acc_holder_name', 'acc_holder_name', 'holder_name'],
    ['Valuta conto bancario', 'currency_id', 'currency_id', 'non presente'],
    ['Dipendente - sesso', 'gender', 'sex (male/female/other)', 'sex (male/female)'],
    ['Dipendente - tipo', 'employee_type (elenco)', 'employee_type (elenco)', 'employee_type_id (anagrafica)'],
    ['Dipendente - n. previdenziale', 'ssnid', 'ssnid', 'non presente'],
    ['Dipendente - conto bancario', 'bank_account_id (uno)', 'bank_account_ids (più conti)', 'bank_account_ids (più conti)'],
    ['Dipendente - versioni', 'non presente', 'date_version', 'date_version'],
    ['Budget analitico', 'budget.analytic', 'invariato', 'invariato (confermato da modulo demo 20.0)'],
    ['Budget finanziario', 'account.report.budget', 'invariato', 'da riconfermare con import di prova'],
])
p('Per usare i template su un database 18 o 19 rinominare le colonne come da tabella (o adattare genera_template.py e rigenerare).')

# ---------------------------------------------------------------- 10 troubleshooting
doc.add_heading('10. Errori frequenti', 1)
table(['Messaggio / sintomo', 'Causa', 'Soluzione'], [
    ['No matching record found for ... in field ...', 'Nome o ID esterno del record collegato inesistente o scritto diversamente',
     'Copiare il nome esatto da Odoo; per /id verificare che il record sia stato importato prima'],
    ['The VAT number ... does not seem to be valid', 'P.IVA errata o senza prefisso paese', 'Correggere il dato (IT + 11 cifre)'],
    ['Duplicati dopo un secondo import', 'Colonna id vuota o cambiata', 'Usare sempre gli stessi id; eliminare i duplicati prima di reimportare'],
    ['CAP / codici senza zeri iniziali', 'Colonna formattata come numero', 'Formattare come Testo e reinserire'],
    ['Il fornitore non compare nel menu Fornitori', 'supplier_rank vuoto', 'Reimportare con supplier_rank = 1'],
    ['Budget analitico senza consuntivo', 'Budget in Bozza o conto analitico nella colonna di piano sbagliata', 'Confermare il budget; verificare x_planN_id'],
    ['Conto receivable/payable: errore riconciliazione', 'reconcile = FALSE', 'Impostare TRUE'],
])

doc.add_heading('11. Checklist finale', 1)
for s in ['Righe di esempio (gialle) eliminate', 'Test di import senza errori su database di staging',
          'Controllo a campione dei record importati', 'File definitivi archiviati come evidenza di migrazione',
          'File con dati personali cancellati dalle postazioni a fine progetto',
          'Approvazione formale del caricamento in produzione (change management)']:
    doc.add_paragraph('☐ ' + s)

out = os.path.join(HERE, 'Manuale_Template_Import_Odoo.docx')
doc.save(out)
print('creato', out)
