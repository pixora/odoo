#!/usr/bin/env python3
"""Converte il foglio 'Matrice_preparazione' (conto x mese) del file
06_Budget_Finanziario.xlsx nel formato importabile del foglio 'Dati'.

Uso:
    python3 budget_matrice_to_import.py 06_Budget_Finanziario.xlsx "Budget 2026" BUDFIN_2026

Crea <file>_IMPORT.xlsx con un solo foglio 'Dati' pronto per Odoo.
Le celle vuote o a zero vengono saltate.
"""
import sys
from datetime import date, datetime

from openpyxl import Workbook, load_workbook


def main(src, budget_name, external_id):
    wb = load_workbook(src, data_only=True)
    ws = wb['Matrice_preparazione']
    header = [c.value for c in ws[2]]  # riga 2: code, nome, mesi...
    months = header[2:]
    out = Workbook()
    dati = out.active
    dati.title = 'Dati'
    dati.append(['id', 'name', 'item_ids/account_id', 'item_ids/date', 'item_ids/amount'])
    first = True
    for row in ws.iter_rows(min_row=3, values_only=True):
        code = row[0]
        if code in (None, ''):
            continue
        for month, amount in zip(months, row[2:]):
            if not amount:
                continue
            if isinstance(month, datetime):
                month = month.date()
            if isinstance(month, date):
                month = month.isoformat()
            head = [external_id, budget_name] if first else ['', '']
            dati.append(head + [str(code), str(month), float(amount)])
            first = False
    dest = src.rsplit('.', 1)[0] + '_IMPORT.xlsx'
    out.save(dest)
    print(f'Creato {dest} con {dati.max_row - 1} righe di budget')


if __name__ == '__main__':
    if len(sys.argv) != 4:
        sys.exit(__doc__)
    main(*sys.argv[1:])
